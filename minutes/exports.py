"""Self-contained document generation: no Office, browser or cloud renderer."""

from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

from minutes.models import Meeting, timestamp


def action_rows(meeting: Meeting) -> list[list[str]]:
    if not meeting.extraction:
        return []
    return [
        [
            item.responsible or "Не указан / Белгісіз",
            item.task,
            item.display_deadline(),
            item.display_status(),
        ]
        for item in meeting.extraction.action_items
    ]


def transcript_lines(meeting: Meeting):
    for segment in meeting.transcript.segments:
        name = meeting.speaker_names.get(segment.speaker, segment.speaker) or "UNKNOWN"
        yield f"[{segment.id}] {timestamp(segment.start)}–{timestamp(segment.end)} {name}: {segment.text}"


def export_docx(meeting: Meeting) -> bytes:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "DejaVu Sans"
    normal.font.size = Pt(10)
    document.add_heading(meeting.title, 0)
    document.add_paragraph(f"Дата / Күні: {meeting.meeting_date}")
    document.add_paragraph("AI draft / Черновик: проверьте имена, задачи и сроки по записи.")
    if meeting.speaker_names:
        document.add_heading("Участники / Қатысушылар", level=1)
        for label, name in meeting.speaker_names.items():
            document.add_paragraph(f"{label}: {name}")
    document.add_heading("Краткое содержание / Қысқаша мазмұны", level=1)
    document.add_paragraph(meeting.extraction.summary if meeting.extraction else "Не сформировано")
    document.add_heading("Задачи / Тапсырмалар", level=1)
    rows = action_rows(meeting)
    if rows:
        table = document.add_table(rows=1, cols=4)
        table.style = "Light Shading Accent 1"
        for cell, text in zip(table.rows[0].cells, ["Ответственный", "Задача", "Срок", "Статус"]):
            cell.text = text
        for row in rows:
            for cell, text in zip(table.add_row().cells, row):
                cell.text = text
        document.add_heading("Основания / Дереккөздер", level=2)
        for i, item in enumerate(meeting.extraction.action_items, 1):
            document.add_paragraph(
                f"{i}. Сегменты {item.evidence_segment_ids}: {item.evidence_quote}"
            )
    else:
        document.add_paragraph("Явные задачи не найдены / Нақты тапсырмалар табылмады")
    document.add_heading("Транскрипт / Транскрипция", level=1)
    for line in transcript_lines(meeting):
        document.add_paragraph(line)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def export_pdf(meeting: Meeting) -> bytes:
    fonts = Path(__file__).resolve().parent.parent / "assets" / "fonts"
    # Register bundled Unicode fonts rather than depending on host fonts or a renderer.
    if "AlemSans" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("AlemSans", str(fonts / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("AlemSansBold", str(fonts / "DejaVuSans-Bold.ttf")))
    body = ParagraphStyle("body", fontName="AlemSans", fontSize=9, leading=13, spaceAfter=6)
    heading = ParagraphStyle(
        "heading", parent=body, fontName="AlemSansBold", fontSize=13, leading=18, spaceBefore=12
    )
    title = ParagraphStyle("title", parent=heading, fontSize=20, leading=25)

    def paragraph(text: str, style=body):
        return Paragraph(escape(text).replace("\n", "<br/>"), style)

    parts = [
        paragraph(meeting.title, title),
        paragraph(f"Дата / Күні: {meeting.meeting_date}"),
        paragraph("AI draft / Черновик: проверьте имена, задачи и сроки по записи."),
    ]
    if meeting.speaker_names:
        parts.append(paragraph("Участники / Қатысушылар", heading))
        parts.extend(paragraph(f"{label}: {name}") for label, name in meeting.speaker_names.items())
    parts += [
        paragraph("Краткое содержание / Қысқаша мазмұны", heading),
        paragraph(meeting.extraction.summary if meeting.extraction else "Не сформировано"),
        paragraph("Задачи / Тапсырмалар", heading),
    ]
    rows = action_rows(meeting)
    if rows:
        table = LongTable(
            [
                [paragraph(v) for v in row]
                for row in [["Ответственный", "Задача", "Срок", "Статус"], *rows]
            ],
            colWidths=[90, 245, 90, 90],
            repeatRows=1,
            splitByRow=1,
            splitInRow=1,
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDEFEA")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C5D3CF")),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        parts += [table, Spacer(1, 10), paragraph("Основания / Дереккөздер", heading)]
        parts.extend(
            paragraph(f"{i}. Сегменты {a.evidence_segment_ids}: {a.evidence_quote}")
            for i, a in enumerate(meeting.extraction.action_items, 1)
        )
    else:
        parts.append(paragraph("Явные задачи не найдены / Нақты тапсырмалар табылмады"))
    parts.append(paragraph("Транскрипт / Транскрипция", heading))
    parts.extend(paragraph(line) for line in transcript_lines(meeting))
    stream = BytesIO()
    document = SimpleDocTemplate(
        stream,
        pagesize=A4,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40,
        title=meeting.title,
        author="Alem Minutes",
    )
    document.build(parts)
    return stream.getvalue()
